"""Tests for universal file conversion across DOCX, PPTX, PDF, Images, Text, and Data."""
from __future__ import annotations

import tempfile
from pathlib import Path
import pytest
from PIL import Image
from docx import Document
from pptx import Presentation

from actions.file_controller import convert_file, _resolve_convert_source
from actions.file_processor import file_processor


def test_resolve_convert_source_direct(tmp_path: Path):
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Sample", 0)
    doc.save(doc_path)

    found = _resolve_convert_source(path=str(tmp_path), name="sample.docx")
    assert found == doc_path

    # Swapped path as filename
    found_swapped = _resolve_convert_source(path=str(doc_path), name="")
    assert found_swapped == doc_path


def test_convert_docx_to_pdf_and_txt(tmp_path: Path):
    doc_path = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("Executive Report", 0)
    doc.add_paragraph("This is an important company report.")
    doc.save(doc_path)

    # DOCX -> TXT
    res_txt = convert_file(path=str(tmp_path), name="report.docx", to="txt")
    assert "Successfully converted" in res_txt or "Converted" in res_txt
    txt_out = tmp_path / "report.txt"
    assert txt_out.exists()
    assert "Executive Report" in txt_out.read_text(encoding="utf-8")

    # DOCX -> PDF
    res_pdf = convert_file(path=str(tmp_path), name="report.docx", to="pdf")
    assert "Successfully converted" in res_pdf
    pdf_out = tmp_path / "report.pdf"
    assert pdf_out.exists()
    assert pdf_out.stat().st_size > 500


def test_convert_pptx_to_pdf_and_txt_and_docx(tmp_path: Path):
    ppt_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Strategic Roadmap"
    prs.save(ppt_path)

    # PPTX -> TXT
    res_txt = convert_file(path=str(tmp_path), name="deck.pptx", to="txt")
    assert "Successfully converted" in res_txt
    txt_out = tmp_path / "deck.txt"
    assert txt_out.exists()
    assert "Strategic Roadmap" in txt_out.read_text(encoding="utf-8")

    # PPTX -> DOCX
    res_docx = convert_file(path=str(tmp_path), name="deck.pptx", to="docx")
    assert "Successfully converted" in res_docx
    docx_out = tmp_path / "deck.docx"
    assert docx_out.exists()
    assert docx_out.stat().st_size > 1000

    # PPTX -> PDF
    res_pdf = convert_file(path=str(tmp_path), name="deck.pptx", to="pdf")
    assert "Successfully converted" in res_pdf
    pdf_out = tmp_path / "deck.pdf"
    assert pdf_out.exists()
    assert pdf_out.stat().st_size > 500


def test_convert_pdf_to_docx_and_txt(tmp_path: Path):
    # Create sample PDF using reportlab
    pdf_path = tmp_path / "article.pdf"
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(pdf_path))
    c.drawString(100, 700, "Quantum Computing Breakthroughs")
    c.drawString(100, 680, "Recent algorithms have achieved linear scaling.")
    c.save()

    # PDF -> TXT
    res_txt = convert_file(path=str(tmp_path), name="article.pdf", to="txt")
    assert "Successfully converted" in res_txt
    txt_out = tmp_path / "article.txt"
    assert txt_out.exists()
    assert "Quantum Computing" in txt_out.read_text(encoding="utf-8")

    # PDF -> DOCX
    res_docx = convert_file(path=str(tmp_path), name="article.pdf", to="docx")
    assert "Successfully converted" in res_docx
    docx_out = tmp_path / "article.docx"
    assert docx_out.exists()
    assert docx_out.stat().st_size > 1000


def test_convert_image_to_pdf_and_format_change(tmp_path: Path):
    img_path = tmp_path / "photo.png"
    im = Image.new("RGBA", (300, 200), color=(50, 100, 200, 255))
    im.save(img_path, "PNG")

    # PNG -> PDF
    res_pdf = convert_file(path=str(tmp_path), name="photo.png", to="pdf")
    assert "Successfully converted" in res_pdf
    pdf_out = tmp_path / "photo.pdf"
    assert pdf_out.exists()
    assert pdf_out.stat().st_size > 200

    # PNG -> JPG
    res_jpg = convert_file(path=str(tmp_path), name="photo.png", to="jpg")
    assert "Successfully converted" in res_jpg
    jpg_out = tmp_path / "photo.jpg"
    assert jpg_out.exists()
    assert jpg_out.stat().st_size > 200


def test_convert_text_markdown_to_pdf_and_docx(tmp_path: Path):
    md_path = tmp_path / "notes.md"
    md_path.write_text("# Project Guidelines\n\n- Write unit tests\n- Clean architecture\n", encoding="utf-8")

    # MD -> PDF
    res_pdf = convert_file(path=str(tmp_path), name="notes.md", to="pdf")
    assert "Successfully converted" in res_pdf
    pdf_out = tmp_path / "notes.pdf"
    assert pdf_out.exists()

    # MD -> DOCX
    res_docx = convert_file(path=str(tmp_path), name="notes.md", to="docx")
    assert "Successfully converted" in res_docx
    docx_out = tmp_path / "notes.docx"
    assert docx_out.exists()


def test_convert_csv_to_excel_and_json(tmp_path: Path):
    csv_path = tmp_path / "metrics.csv"
    csv_path.write_text("Year,Revenue,Growth\n2024,100,0.15\n2025,125,0.25\n", encoding="utf-8")

    # CSV -> XLSX
    res_xlsx = convert_file(path=str(tmp_path), name="metrics.csv", to="xlsx")
    assert "Successfully converted" in res_xlsx
    xlsx_out = tmp_path / "metrics.xlsx"
    assert xlsx_out.exists()

    # CSV -> JSON
    res_json = convert_file(path=str(tmp_path), name="metrics.csv", to="json")
    assert "Successfully converted" in res_json
    json_out = tmp_path / "metrics.json"
    assert json_out.exists()
    assert "2024" in json_out.read_text(encoding="utf-8")


def test_file_processor_routing_for_conversion(tmp_path: Path):
    doc_path = tmp_path / "summary.docx"
    doc = Document()
    doc.add_heading("Quarterly Summary", 0)
    doc.save(doc_path)

    res = file_processor({"file_path": str(doc_path), "action": "to_pdf"})
    assert "Successfully converted" in res
    assert (tmp_path / "summary.pdf").exists()
