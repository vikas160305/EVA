"""Create finished documents, presentations, PDFs, and source code from a request."""
from __future__ import annotations

import json
import html
import re
import tempfile
import time
import os
from pathlib import Path

import requests

from actions.file_controller import DocumentService


BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = Path.home() / "Documents" / "JARVIS Created"
CONTENT_MODELS = [
    "gemini-3.5-flash",
    "gemini-2.5-flash",
    "gemini-flash-latest",
    "gemini-3.8-flash",
    "gemini-flash-lite-latest",
    "gemini-3.6-flash",
]
CONTENT_MODEL = CONTENT_MODELS[0]

_CODE_EXTENSIONS = {
    "python": ".py", "javascript": ".js", "typescript": ".ts",
    "html": ".html", "css": ".css", "java": ".java", "c": ".c",
    "cpp": ".cpp", "c++": ".cpp", "csharp": ".cs", "go": ".go",
    "rust": ".rs", "php": ".php", "sql": ".sql", "powershell": ".ps1",
}
_TEXT_FORMATS = {"txt", "md", "html", "json"}


def _api_key() -> str:
    try:
        from core.llm_client import get_gemini_api_key
        return get_gemini_api_key()
    except Exception:
        with open(BASE_DIR / "config" / "api_keys.json", encoding="utf-8") as f:
            return json.load(f)["gemini_api_key"]


def _safe_name(value: str, fallback: str = "jarvis_document") -> str:
    value = re.sub(r"[^A-Za-z0-9 _-]", "", value).strip().replace(" ", "_")
    return value[:60] or fallback


def _available_path(destination: Path, stem: str, extension: str) -> Path:
    """Avoid replacing a previous user-created file."""
    candidate = destination / f"{stem}.{extension}"
    if not candidate.exists():
        return candidate
    for index in range(2, 1000):
        candidate = destination / f"{stem}_{index}.{extension}"
        if not candidate.exists():
            return candidate
    raise RuntimeError("Too many files with the same name in the output folder.")


def _research(topic: str) -> str:
    """Use the project's grounded web search when freshness is requested."""
    try:
        from actions.web_search import web_search
        result = web_search({"query": topic, "mode": "research"})
        return result[:12000]
    except Exception as exc:
        return f"No web research available ({exc}). Use general knowledge and say so."


def _generate(prompt: str) -> str:
    from google import genai
    try:
        from core.llm_client import get_gemini_api_key, rotate_gemini_api_key
        current_key = get_gemini_api_key()
    except Exception:
        current_key = _api_key()
        rotate_gemini_api_key = lambda *a, **kw: current_key

    client = genai.Client(api_key=current_key)
    
    candidate_models = list(CONTENT_MODELS)
    last_exc = None
    for model_name in candidate_models:
        for attempt in range(2):
            try:
                response = client.models.generate_content(model=model_name, contents=prompt)
                text = (response.text or "").strip()
                if text:
                    return text
            except Exception as exc:
                last_exc = exc
                err_msg = str(exc)
                if "429" in err_msg or "RESOURCE_EXHAUSTED" in err_msg:
                    current_key = rotate_gemini_api_key(current_key, reason="rate_limit")
                    client = genai.Client(api_key=current_key)
                    time.sleep(0.5)
                    continue
                if "503" in err_msg:
                    time.sleep(1.0)
                    continue
                # For 404 or unsupported models, break to next candidate model
                break
    if last_exc:
        raise last_exc
    return ""


def _json_plan(request: str, format_name: str, research: str, count: int) -> dict:
    prompt = f'''You are an expert content strategist and professional author. Create an authoritative, factual, well-structured plan for a {format_name}.
User request: {request}
Requested length/count: {count}
Research notes (may be empty): {research}

Guidelines:
1. Deliver comprehensive, high-quality, professional content with high factual accuracy.
2. For presentations (PPTX): Create exactly {count} content sections.
   - Each section needs 3 to 5 detailed, insightful bullet points containing concrete facts, data, explanations, and actionable takeaways (avoid vague one-liners).
   - "body": A concise 1-2 sentence executive summary or presenter note for the slide.
   - "image_query": A specific, descriptive 3-5 word image search query for high-res web photos (e.g. "James Webb Space Telescope nebula", "modern server room datacenter", "renewable solar energy farm").
3. For documents (DOCX, PDF): Create 4 to 8 rich, well-organized sections.
   - "body": Substantive, thorough multi-sentence paragraphs (100-250 words per section) explaining core concepts, context, and findings.
   - "points": Bulleted key takeaways or specific metrics.
   - "image_query": A specific, descriptive image search query for section illustration.

Return ONLY valid JSON with this exact shape:
{{"title":"...","subtitle":"...","sections":[{{"title":"...","points":["..."],"body":"...","image_query":"..."}}],"sources":["..."]}}
Do not invent citations. Include source URLs only when present in research.'''
    raw = _generate(prompt)
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.I)
    try:
        plan = json.loads(raw)
        if isinstance(plan.get("sections"), list) and plan["sections"]:
            plan.setdefault("title", _safe_name(request, "Untitled"))
            plan.setdefault("subtitle", "Created by JARVIS")
            plan.setdefault("sources", [])
            for section in plan["sections"]:
                if isinstance(section, dict):
                    section.setdefault("title", "Section")
                    section.setdefault("points", [])
                    section.setdefault("body", "")
                    section.setdefault("image_query", section["title"])
            return plan
    except json.JSONDecodeError:
        pass
    return {"title": _safe_name(request, "Untitled"), "subtitle": "Created by JARVIS", "sections": [
        {"title": "Overview", "points": [raw[:500] or request], "body": raw[:2500], "image_query": request}
    ], "sources": []}


def _fit_sections(plan: dict, count: int) -> list[dict]:
    """Guarantee a requested slide count with diverse, progressive sequential topics."""
    sections = [item for item in plan.get("sections", []) if isinstance(item, dict)]
    if not sections:
        sections = [{"title": "Overview", "points": [plan.get("title", "Overview")], "body": "", "image_query": plan.get("title", "")}]

    subtopic_templates = [
        ("Key Architecture & Foundations", "Core architectural pillars and underlying mechanisms.", "architecture system overview"),
        ("Deep Technical Analysis", "Detailed analysis of core performance characteristics and metrics.", "data analytics dashboard chart"),
        ("Real-World Use Cases & Applications", "Practical implementations, industry workflows, and case studies.", "technology workflow application"),
        ("Key Challenges & Mitigations", "Operational obstacles, edge cases, and effective mitigation strategies.", "solution engineering puzzle"),
        ("Strategic Recommendations", "Actionable best practices and recommendations for optimal outcomes.", "strategy checklist goal"),
        ("Future Roadmap & Innovations", "Emerging innovations, scalability roadmap, and upcoming horizons.", "future technology innovation"),
        ("Summary & Key Takeaways", "Comprehensive synthesis of critical insights and conclusion.", "summary conclusion milestone"),
    ]
    template_idx = 0
    main_topic = plan.get("title", "Topic")
    while len(sections) < count:
        title_tmpl, body_tmpl, query_tmpl = subtopic_templates[template_idx % len(subtopic_templates)]
        template_idx += 1
        sections.append({
            "title": f"{title_tmpl}",
            "points": [
                f"In-depth analysis of {main_topic.lower()} with respect to {title_tmpl.lower()}.",
                f"Identified critical success factors and strategic considerations for execution.",
                f"Actionable methodology aligned with current industry best practices."
            ],
            "body": f"{body_tmpl} Focuses on high-impact dimensions of {main_topic}.",
            "image_query": f"{main_topic} {query_tmpl}",
        })
    return sections[:count]


def _wikimedia_image(query: str, folder: Path) -> Path | None:
    """Fetch one openly hosted thumbnail for a slide from Wikimedia Commons."""
    try:
        response = requests.get(
            "https://commons.wikimedia.org/w/api.php",
            params={"action": "query", "generator": "search", "gsrsearch": query,
                    "gsrnamespace": 6, "gsrlimit": 1, "prop": "imageinfo",
                    "iiprop": "url", "iiurlwidth": 1200, "format": "json"}, timeout=12,
        )
        pages = response.json().get("query", {}).get("pages", {}).values()
        info = next(iter(pages), {}).get("imageinfo", [{}])[0]
        url = info.get("thumburl") or info.get("url")
        if not url:
            return None
        data = requests.get(url, timeout=18).content
        target = folder / f"image_{_safe_name(query, 'slide')}.jpg"
        target.write_bytes(data)
        return target
    except Exception:
        return None


def _fetch_relevant_image(query: str, folder: Path, used_urls: set[str] | None = None) -> Path | None:
    """Search the internet for high-quality, relevant images using DuckDuckGo, Bing, Google, and Wikimedia.
    Validates dimensions and converts to standard RGB mode for perfect embedding."""
    if not query or not query.strip():
        return None
    cleaned_query = re.sub(r"[^\w\s-]", " ", query).strip()
    if not cleaned_query:
        return None

    seen = used_urls if used_urls is not None else set()
    candidate_urls: list[str] = []

    # 1. DuckDuckGo Image Search
    try:
        from actions.image_downloader import _duckduckgo_images
        candidate_urls.extend(_duckduckgo_images(cleaned_query))
    except Exception:
        pass

    # 2. Bing Image Search
    if len(candidate_urls) < 6:
        try:
            from actions.image_downloader import _bing_images
            candidate_urls.extend(_bing_images(cleaned_query))
        except Exception:
            pass

    # 3. Google Image Search
    if len(candidate_urls) < 6:
        try:
            from actions.image_downloader import _google_images
            candidate_urls.extend(_google_images(cleaned_query))
        except Exception:
            pass

    # 4. Wikimedia Commons fallback
    if len(candidate_urls) < 3:
        try:
            resp = requests.get(
                "https://commons.wikimedia.org/w/api.php",
                params={"action": "query", "generator": "search", "gsrsearch": cleaned_query,
                        "gsrnamespace": 6, "gsrlimit": 3, "prop": "imageinfo",
                        "iiprop": "url", "iiurlwidth": 1200, "format": "json"},
                headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"},
                timeout=8,
            )
            pages = resp.json().get("query", {}).get("pages", {}).values()
            for p in pages:
                info = p.get("imageinfo", [{}])[0]
                u = info.get("thumburl") or info.get("url")
                if u:
                    candidate_urls.append(u)
        except Exception:
            pass

    from io import BytesIO
    from PIL import Image

    for url in candidate_urls:
        if not url or url in seen:
            continue
        seen.add(url)
        try:
            res = requests.get(
                url,
                headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"},
                timeout=(5, 12),
            )
            if res.status_code != 200 or len(res.content) < 5000:
                continue
            img = Image.open(BytesIO(res.content))
            width, height = img.size
            if width < 250 or height < 180:
                continue
            if img.format not in ("JPEG", "PNG", "WEBP", "BMP"):
                continue
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")

            slug = _safe_name(cleaned_query, "visual")[:25]
            out_file = folder / f"img_{len(seen)}_{slug}.jpg"
            img.save(out_file, "JPEG", quality=90)
            if out_file.exists() and out_file.stat().st_size > 0:
                return out_file
        except Exception:
            continue

    # Final fallback to standard wikimedia
    return _wikimedia_image(cleaned_query, folder)


def _create_docx(plan: dict, output: Path, include_images: bool = False) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    from PIL import Image

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        footer = s.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run(f"EVA AI  |  Executive Document  |  {time.strftime('%B %Y')}")
        r_ft.font.name = "Calibri"
        r_ft.font.size = Pt(9)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(30, 41, 59)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run(plan.get("title", "Document"))
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy #1E3A8A
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle / Metadata
    subtitle = plan.get("subtitle") or "Prepared by EVA AI"
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = p_sub.add_run(f"{subtitle}  |  {time.strftime('%B %d, %Y')}")
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)
    p_sub.paragraph_format.space_after = Pt(14)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    p_div_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="CBD5E1"/>'
                             r'</w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_div_border)

    with tempfile.TemporaryDirectory() as temp:
        image_dir = Path(temp)
        used_urls: set[str] = set()

        for item in plan.get("sections", []):
            h = doc.add_heading(item.get("title", "Section"), level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.color.rgb = RGBColor(30, 58, 138)
                r.font.size = Pt(16)

            body = (item.get("body") or "").strip()
            if body:
                p_body = doc.add_paragraph(body)
                p_body.paragraph_format.space_after = Pt(8)
                p_body.paragraph_format.line_spacing = 1.15

            for point in item.get("points", []):
                p_pt = doc.add_paragraph(str(point), style="List Bullet")
                p_pt.paragraph_format.space_after = Pt(4)

            # Insert relevant web image if requested
            if include_images:
                query = item.get("image_query") or item.get("title", "")
                img_path = _fetch_relevant_image(query, image_dir, used_urls)
                if img_path and img_path.exists():
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(8)
                        p_img.paragraph_format.space_after = Pt(4)
                        doc.add_picture(str(img_path), width=Inches(5.0))

                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_after = Pt(14)
                        r_cap = p_cap.add_run(f"Figure: {item.get('title', 'Illustration')}")
                        r_cap.font.size = Pt(9)
                        r_cap.font.italic = True
                        r_cap.font.color.rgb = RGBColor(100, 116, 139)
                    except Exception:
                        pass

        if plan.get("sources"):
            h_src = doc.add_heading("References & Sources", level=1)
            h_src.paragraph_format.space_before = Pt(18)
            h_src.paragraph_format.space_after = Pt(6)
            for s in plan["sources"]:
                p_s = doc.add_paragraph(str(s), style="List Bullet")
                p_s.paragraph_format.space_after = Pt(3)

    doc.save(output)


def _create_pdf(plan: dict, output: Path, include_images: bool = False) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, HRFlowable, Image as RLImage
    from PIL import Image as PILImage

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=24,
        leading=28,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=6,
    )
    sub_style = ParagraphStyle(
        "DocSub",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=12,
    )
    h1_style = ParagraphStyle(
        "DocH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=8,
    )
    bullet_style = ParagraphStyle(
        "DocBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=4,
    )
    caption_style = ParagraphStyle(
        "DocCaption",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        alignment=1,  # Center
        spaceAfter=12,
    )

    story = [
        Paragraph(plan.get("title", "Document"), title_style),
        Paragraph(f"{plan.get('subtitle', 'Prepared by EVA AI')}  •  {time.strftime('%B %d, %Y')}", sub_style),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=14),
    ]

    with tempfile.TemporaryDirectory() as temp:
        image_dir = Path(temp)
        used_urls: set[str] = set()

        for item in plan.get("sections", []):
            story.append(Paragraph(item.get("title", "Section"), h1_style))
            body = (item.get("body") or "").strip()
            if body:
                story.append(Paragraph(body, body_style))

            points = item.get("points", [])
            if points:
                items = [ListItem(Paragraph(str(p), bullet_style), leftIndent=15, bulletOffsetY=-1) for p in points]
                story.append(ListFlowable(items, bulletType="bullet", start="square", leftIndent=10))
                story.append(Spacer(1, 0.08 * inch))

            if include_images:
                query = item.get("image_query") or item.get("title", "")
                img_path = _fetch_relevant_image(query, image_dir, used_urls)
                if img_path and img_path.exists():
                    try:
                        with PILImage.open(img_path) as pim:
                            pw, ph = pim.size
                        aspect = pw / max(1, ph)
                        max_w, max_h = 440, 240
                        target_w = max_w
                        target_h = max_w / aspect
                        if target_h > max_h:
                            target_h = max_h
                            target_w = max_h * aspect

                        story.append(Spacer(1, 0.08 * inch))
                        story.append(RLImage(str(img_path), width=target_w, height=target_h))
                        story.append(Paragraph(f"Figure: {item.get('title', 'Illustration')}", caption_style))
                    except Exception:
                        pass

        if plan.get("sources"):
            story.append(Paragraph("References & Sources", h1_style))
            items = [ListItem(Paragraph(str(s), bullet_style), leftIndent=15) for s in plan["sources"]]
            story.append(ListFlowable(items, bulletType="bullet", leftIndent=10))

    from reportlab.pdfgen import canvas as rl_canvas

    class NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_decorations(num_pages)
                super().showPage()
            super().save()

        def draw_page_decorations(self, page_count):
            self.saveState()
            self.setFont("Helvetica", 9)
            self.setFillColor(colors.HexColor("#94A3B8"))
            if self._pageNumber > 1:
                self.setStrokeColor(colors.HexColor("#E2E8F0"))
                self.setLineWidth(0.5)
                self.line(54, 11 * 72 - 36, 8.5 * 72 - 54, 11 * 72 - 36)
                self.drawString(54, 11 * 72 - 30, "EVA AI — Executive Document")

            self.setStrokeColor(colors.HexColor("#E2E8F0"))
            self.setLineWidth(0.5)
            self.line(54, 45, 8.5 * 72 - 54, 45)
            self.drawString(54, 32, "Prepared by EVA AI")
            page_str = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(8.5 * 72 - 54, 32, page_str)
            self.restoreState()

    doc = SimpleDocTemplate(str(output), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    doc.build(story, canvasmaker=NumberedCanvas)


def _create_pptx(plan: dict, output: Path, total_slides: int, include_images: bool = False) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # Slide 1: Executive Title Slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42)  # #0F172A Slate Navy
    bg.line.fill.background()

    accent = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(1.8), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(37, 99, 235)  # Royal Blue
    accent.line.fill.background()

    title_box = title_slide.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(11.0), Inches(2.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = plan.get("title", "Presentation")
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    sub_box = title_slide.shapes.add_textbox(Inches(1.2), Inches(4.7), Inches(10.5), Inches(1.2))
    stf = sub_box.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.text = plan.get("subtitle") or "Created by EVA AI"
    sp.font.size = Pt(20)
    sp.font.color.rgb = RGBColor(203, 213, 225)

    foot_box = title_slide.shapes.add_textbox(Inches(1.2), Inches(6.3), Inches(8.0), Inches(0.5))
    fp = foot_box.text_frame.paragraphs[0]
    fp.text = "Generated by EVA AI • Executive Deck"
    fp.font.size = Pt(11)
    fp.font.color.rgb = RGBColor(148, 163, 184)

    with tempfile.TemporaryDirectory() as temp:
        image_dir = Path(temp)
        used_urls: set[str] = set()
        sections = _fit_sections(plan, max(1, total_slides - 1))

        for index, item in enumerate(sections, 1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])

            # Top accent bar
            accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(1.5), Inches(0.06))
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = RGBColor(37, 99, 235)
            accent_bar.line.fill.background()

            # Slide Title
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(10.0), Inches(0.65))
            ttf = title_box.text_frame
            ttf.word_wrap = True
            tp = ttf.paragraphs[0]
            tp.text = item.get("title", f"Section {index}")
            tp.font.size = Pt(26)
            tp.font.bold = True
            tp.font.color.rgb = RGBColor(15, 23, 42)

            # Slide index badge
            cat_box = slide.shapes.add_textbox(Inches(10.0), Inches(0.5), Inches(2.5), Inches(0.4))
            cp = cat_box.text_frame.paragraphs[0]
            cp.text = f"{index:02d} / {len(sections):02d}"
            cp.alignment = PP_ALIGN.RIGHT
            cp.font.size = Pt(12)
            cp.font.color.rgb = RGBColor(148, 163, 184)

            # Check if image is requested and fetch
            image_path = None
            if include_images:
                query = item.get("image_query") or item.get("title", "")
                image_path = _fetch_relevant_image(query, image_dir, used_urls)

            # Layout logic: Two-column with Image vs Full-Width
            if image_path and image_path.exists():
                text_left = Inches(0.8)
                text_top = Inches(1.45)
                text_width = Inches(6.8)
                text_height = Inches(5.4)

                box_left = Inches(7.9)
                box_top = Inches(1.45)
                box_width = Inches(4.6)
                box_height = Inches(5.2)

                try:
                    with Image.open(image_path) as im:
                        im_w, im_h = im.size
                    aspect = im_w / max(1, im_h)
                    target_w = box_width
                    target_h = box_width / aspect
                    if target_h > box_height:
                        target_h = box_height
                        target_w = box_height * aspect

                    img_x = box_left + (box_width - target_w) / 2
                    img_y = box_top + (box_height - target_h) / 2

                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_left, box_top, box_width, box_height)
                    card.fill.solid()
                    card.fill.fore_color.rgb = RGBColor(248, 250, 252)
                    card.line.color.rgb = RGBColor(226, 232, 240)

                    slide.shapes.add_picture(str(image_path), img_x, img_y, width=target_w, height=target_h)
                except Exception:
                    text_width = Inches(11.7)
            else:
                text_left = Inches(0.8)
                text_top = Inches(1.45)
                text_width = Inches(11.7)
                text_height = Inches(5.4)

            # Content Text Box
            box = slide.shapes.add_textbox(text_left, text_top, text_width, text_height)
            frame = box.text_frame
            frame.word_wrap = True

            first_p = True
            body_text = (item.get("body") or "").strip()
            if body_text and len(body_text) > 10:
                p = frame.paragraphs[0]
                first_p = False
                p.text = body_text
                p.font.size = Pt(14)
                p.font.italic = True
                p.font.color.rgb = RGBColor(71, 85, 105)
                p.space_after = Pt(14)

            points = item.get("points", [])
            for point in points[:5]:
                point_text = str(point).strip()
                if not point_text:
                    continue
                if first_p:
                    p = frame.paragraphs[0]
                    first_p = False
                else:
                    p = frame.add_paragraph()
                p.text = f"•  {point_text}"
                p.level = 0
                p.font.size = Pt(16 if (image_path and image_path.exists()) else 17)
                p.font.color.rgb = RGBColor(30, 41, 59)
                p.space_after = Pt(12)

    prs.save(output)


def _create_text(plan: dict, output: Path, format_name: str) -> None:
    if format_name == "json":
        output.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        return
    lines = [plan["title"], "", plan.get("subtitle", ""), ""]
    for item in plan["sections"]:
        heading = item.get("title", "Section")
        if format_name == "html":
            lines += [f"<h2>{html.escape(str(heading))}</h2>", f"<p>{html.escape(str(item.get('body', '')))}</p>", "<ul>"]
            lines += [f"<li>{html.escape(str(point))}</li>" for point in item.get("points", [])]
            lines += ["</ul>"]
        else:
            lines += [f"## {heading}", item.get("body", "")]
            lines += [f"- {point}" for point in item.get("points", [])]
            lines += [""]
    if plan.get("sources"):
        if format_name == "html":
            lines += ["<h2>Sources</h2>", "<ul>"]
            lines += [f"<li>{html.escape(str(source))}</li>" for source in plan["sources"]]
            lines += ["</ul>"]
        else:
            lines += ["## Sources", *[str(source) for source in plan["sources"]]]
    body = "\n".join(lines)
    if format_name == "html":
        output.write_text(
            "<!doctype html><html><head><meta charset=\"utf-8\"><title>"
            + html.escape(str(plan["title"]))
            + f"</title></head><body><h1>{html.escape(str(plan['title']))}</h1>"
            + f"<p>{html.escape(str(plan.get('subtitle', '')))}</p>"
            + body
            + "</body></html>",
            encoding="utf-8",
        )
    else:
        output.write_text(body, encoding="utf-8")


def _create_code(request: str, language: str, output: Path) -> None:
    prompt = f"""Write complete production-ready {language} code for this request: {request}
Return ONLY code, without Markdown fences. Include clear error handling and comments where useful."""
    code = re.sub(r"^```[\w+-]*\s*|\s*```$", "", _generate(prompt).strip(), flags=re.S)
    output.write_text(code, encoding="utf-8")


def content_creator(parameters: dict, response=None, player=None, session_memory=None, speak=None) -> str:
    params = parameters or {}
    request = str(params.get("request") or params.get("description") or "").strip()
    format_name = str(params.get("format", "")).lower().strip(". ")
    requested_code = format_name in {"code", "program", "script"}
    if not request:
        return "Please say what you want to create and the format, for example a PowerPoint, Word document, PDF, or code file."
    if format_name in {"ppt", "powerpoint", "slides"}: format_name = "pptx"
    if format_name in {"word", "document", "doc"}: format_name = "docx"
    if format_name in {"markdown", "text"}: format_name = "md" if format_name == "markdown" else "txt"
    language = str(params.get("language", "python")).lower()
    if requested_code:
        format_name = _CODE_EXTENSIONS.get(language, "." + _safe_name(language, "txt").lower()).lstrip(".")
    supported_code_extensions = {e.lstrip(".") for e in _CODE_EXTENSIONS.values()}
    if format_name not in {"pdf", "docx", "pptx", *_TEXT_FORMATS, *supported_code_extensions} and not requested_code:
        return "Supported creation formats are PDF, DOCX/Word, PPTX/PowerPoint, TXT, Markdown, HTML, JSON, and code languages."
    count = max(1, min(int(params.get("count", params.get("slides", 10))), 30))
    req_low = request.lower()
    default_dest = OUTPUT_DIR
    if (requested_code or format_name in supported_code_extensions) and ("desktop" in req_low or any(term in req_low for term in ("vscode", "vs code", "vls code"))):
        try:
            from actions.code_helper import _get_desktop_dir
            default_dest = _get_desktop_dir()
        except Exception:
            pass
    destination = Path(params.get("destination") or default_dest).expanduser()
    destination.mkdir(parents=True, exist_ok=True)

    service = DocumentService()
    raw_name = params.get('name')
    if not raw_name and (requested_code or (format_name in supported_code_extensions and format_name not in _TEXT_FORMATS)):
        try:
            from actions.code_helper import _infer_slug_from_description
            slug = _infer_slug_from_description(request, fallback="Main")
            raw_name = slug.capitalize() if format_name == "java" else slug
        except Exception:
            raw_name = request
    else:
        raw_name = raw_name or request

    output = service._next_safe_path(destination, _safe_name(raw_name), format_name)
    started_at = time.monotonic()
    if player:
        player.write_log(f"[content_creator] Creating {format_name.upper()} at {output}")
    try:
        tmp = service._atomic_tmp_path(output)
        if requested_code or (format_name in supported_code_extensions and format_name not in _TEXT_FORMATS):
            _create_code(request, language, tmp)
            service._validate_artifact(tmp, format_name)
            os.replace(tmp, output)
        else:
            research = _research(request) if params.get("research", True) else ""
            plan = _json_plan(request, format_name, research, max(1, count - 1) if format_name == "pptx" else count)
            req_words = ("image", "images", "photo", "photos", "picture", "pictures", "visual", "visuals", "diagram", "diagrams", "illustration", "illustrations", "figure", "figures")
            include_images = bool(params.get("images") or params.get("include_images")) or any(w in req_low for w in req_words)
            if format_name == "docx":
                _create_docx(plan, tmp, include_images=include_images)
                service._validate_artifact(tmp, "docx")
                os.replace(tmp, output)
            elif format_name == "pdf":
                _create_pdf(plan, tmp, include_images=include_images)
                service._validate_artifact(tmp, "pdf")
                os.replace(tmp, output)
            elif format_name == "pptx":
                _create_pptx(plan, tmp, count, include_images=include_images)
                service._validate_artifact(tmp, "pptx")
                os.replace(tmp, output)
            else:
                _create_text(plan, tmp, format_name)
                service._validate_artifact(tmp, format_name)
                os.replace(tmp, output)
        if not output.exists() or output.stat().st_size == 0:
            raise RuntimeError("The output file was not written correctly.")
        
        # Auto-reveal created file in File Explorer
        try:
            from actions.media_downloader import _reveal_in_explorer
            _reveal_in_explorer(output)
        except Exception:
            pass

        result = f"Created {format_name.upper()} file: {output} ({time.monotonic() - started_at:.1f}s)"
        if requested_code or (format_name in supported_code_extensions and format_name not in _TEXT_FORMATS):
            if any(term in req_low for term in ("vscode", "vs code", "vls code", "editor", "in code")):
                try:
                    from actions.code_helper import _open_vscode
                    if _open_vscode(output):
                        result += "\n💻 Opened in Visual Studio Code."
                except Exception:
                    pass
        if speak:
            speak(result)
        return result
    except Exception as exc:
        detail = f"{type(exc).__name__}: {exc}"
        print(f"[content_creator] FAILED — {detail}")
        if player:
            player.write_log(f"[content_creator] ERROR: {detail[:220]}")
        return f"Could not create the {format_name.upper()} file: {detail}"
