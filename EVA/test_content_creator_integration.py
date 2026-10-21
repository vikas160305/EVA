"""Integration tests for document and content creation."""
import pytest
from pathlib import Path
from actions.content_creator import content_creator


def test_content_creator_empty_request():
    result = content_creator({})
    assert "Please say what you want to create" in result


def test_content_creator_unsupported_format():
    result = content_creator({"request": "Hello world", "format": "xyz_unsupported"})
    assert "Supported creation formats" in result


def test_content_creator_code_generation(tmp_path: Path):
    result = content_creator({
        "request": "Write a python function to add two numbers",
        "format": "code",
        "language": "python",
        "destination": str(tmp_path)
    })
    assert "Created" in result or "PY file" in result
    py_files = list(tmp_path.glob("*.py"))
    assert len(py_files) >= 1
    content = py_files[0].read_text(encoding="utf-8")
    assert "def " in content


def test_content_creator_txt_generation(tmp_path: Path):
    result = content_creator({
        "request": "A short summary of artificial intelligence",
        "format": "txt",
        "destination": str(tmp_path)
    })
    assert "Created TXT file" in result
    txt_files = list(tmp_path.glob("*.txt"))
    assert len(txt_files) >= 1
    assert txt_files[0].stat().st_size > 0


def test_fit_sections_diverse_topics():
    from actions.content_creator import _fit_sections
    plan = {
        "title": "Autonomous Robotics",
        "sections": [
            {"title": "Introduction to Robotics", "points": ["Perception systems", "Actuators"], "body": "Overview", "image_query": "robot"}
        ]
    }
    fitted = _fit_sections(plan, 5)
    assert len(fitted) == 5
    # Verify titles are diverse and not repeated "continued"
    titles = [s["title"] for s in fitted]
    assert len(set(titles)) == 5
    assert not any("— continued" in t for t in titles)


def test_create_pptx_with_images_and_styling(tmp_path: Path, monkeypatch):
    from PIL import Image
    from pptx import Presentation
    from actions.content_creator import _create_pptx

    # Create dummy test image
    test_img = tmp_path / "test_visual.jpg"
    im = Image.new("RGB", (400, 300), color=(30, 80, 180))
    im.save(test_img, "JPEG")

    monkeypatch.setattr("actions.content_creator._fetch_relevant_image", lambda q, f, u=None: test_img)

    plan = {
        "title": "Space Exploration 2030",
        "subtitle": "Deep Space Roadmap",
        "sections": [
            {
                "title": "Lunar Gateway Station",
                "points": ["Permanent orbit around the Moon", "Habitation and scientific research", "Staging platform for Mars"],
                "body": "The Gateway serves as humanity's outpost in deep space.",
                "image_query": "lunar gateway station space"
            },
            {
                "title": "Artemis Lander Architecture",
                "points": ["Human Landing System", "Reusable orbital propulsion", "Autonomous guidance"],
                "body": "Crew landing capabilities on the South Pole.",
                "image_query": "artemis moon lander"
            }
        ]
    }

    out_pptx = tmp_path / "space_presentation.pptx"
    _create_pptx(plan, out_pptx, total_slides=3, include_images=True)
    assert out_pptx.exists()
    assert out_pptx.stat().st_size > 1000

    prs = Presentation(str(out_pptx))
    assert len(prs.slides) >= 3
    # Check title slide
    assert "Space Exploration" in prs.slides[0].shapes[2].text_frame.text
    # Check content slide has picture
    has_pic = any(shape.shape_type == 13 for shape in prs.slides[1].shapes)  # 13 = MSO_SHAPE_TYPE.PICTURE
    assert has_pic


def test_create_docx_with_images_and_styling(tmp_path: Path, monkeypatch):
    from PIL import Image
    from docx import Document
    from actions.content_creator import _create_docx

    test_img = tmp_path / "test_doc_img.jpg"
    im = Image.new("RGB", (400, 300), color=(40, 120, 200))
    im.save(test_img, "JPEG")

    monkeypatch.setattr("actions.content_creator._fetch_relevant_image", lambda q, f, u=None: test_img)

    plan = {
        "title": "Machine Learning in Production",
        "subtitle": "Engineering Whitepaper",
        "sections": [
            {
                "title": "Model Serving & Latency",
                "points": ["ONNX runtime optimizations", "Batch inference microservices"],
                "body": "Low-latency inference is critical for interactive systems.",
                "image_query": "server inference cluster"
            }
        ],
        "sources": ["https://arxiv.org/abs/example"]
    }

    out_docx = tmp_path / "ml_production.docx"
    _create_docx(plan, out_docx, include_images=True)
    assert out_docx.exists()
    assert out_docx.stat().st_size > 1000

    doc = Document(str(out_docx))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Machine Learning in Production" in all_text
    assert "Model Serving & Latency" in all_text


def test_create_pdf_with_images_and_styling(tmp_path: Path, monkeypatch):
    from PIL import Image
    from actions.content_creator import _create_pdf

    test_img = tmp_path / "test_pdf_img.jpg"
    im = Image.new("RGB", (400, 300), color=(60, 140, 220))
    im.save(test_img, "JPEG")

    monkeypatch.setattr("actions.content_creator._fetch_relevant_image", lambda q, f, u=None: test_img)

    plan = {
        "title": "Renewable Energy Transitions",
        "subtitle": "Global Climate Analysis",
        "sections": [
            {
                "title": "Solar Photovoltaic Efficiency",
                "points": ["Perovskite tandem cells", "Grid integration storage solutions"],
                "body": "Solar energy deployment has expanded exponentially over the last decade.",
                "image_query": "solar photovoltaic farm"
            }
        ],
        "sources": ["International Energy Agency Report 2025"]
    }

    out_pdf = tmp_path / "energy_report.pdf"
    _create_pdf(plan, out_pdf, include_images=True)
    assert out_pdf.exists()
    assert out_pdf.stat().st_size > 1000
